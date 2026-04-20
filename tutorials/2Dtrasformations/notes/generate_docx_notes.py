"""
Generate professionally formatted DOCX files from transformation markdown notes.

Creates:
- One .docx per source .md note
- One combined .docx containing all notes
"""

from __future__ import annotations

from datetime import date
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE_FILES = [
    "01_2d_translation.md",
    "02_2d_scaling.md",
    "03_2d_rotation.md",
    "04_2d_reflection.md",
    "05_2d_shearing.md",
    "06_composing_transformations_detailed.md",
]

COMBINED_OUTPUT = "00_all_2d_transformations_notes.docx"
IMAGE_PATTERN = re.compile(r"^!\[(.*?)\]\((.*?)\)\s*$")


def configure_document(doc: Document) -> None:
    """Apply global style configuration for professional output."""
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    heading1 = doc.styles["Heading 1"]
    heading1.font.name = "Calibri"
    heading1.font.size = Pt(18)
    heading1.font.bold = True
    heading1.font.color.rgb = RGBColor(31, 62, 102)

    heading2 = doc.styles["Heading 2"]
    heading2.font.name = "Calibri"
    heading2.font.size = Pt(14)
    heading2.font.bold = True
    heading2.font.color.rgb = RGBColor(34, 82, 130)

    heading3 = doc.styles["Heading 3"]
    heading3.font.name = "Calibri"
    heading3.font.size = Pt(12)
    heading3.font.bold = True
    heading3.font.color.rgb = RGBColor(50, 90, 135)


def add_page_number(paragraph) -> None:
    """Insert PAGE field into a paragraph."""
    run = paragraph.add_run("Page ")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(100, 100, 100)

    field_run = paragraph.add_run()
    field_run.font.size = Pt(9)
    field_run.font.color.rgb = RGBColor(100, 100, 100)
    field_r = field_run._r

    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "

    fld_char_sep = OxmlElement("w:fldChar")
    fld_char_sep.set(qn("w:fldCharType"), "separate")

    placeholder = OxmlElement("w:t")
    placeholder.text = "1"

    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")

    field_r.append(fld_char_begin)
    field_r.append(instr_text)
    field_r.append(fld_char_sep)
    field_r.append(placeholder)
    field_r.append(fld_char_end)


def add_header_footer(doc: Document, header_text: str) -> None:
    """Add professional header and footer with page numbers."""
    section = doc.sections[0]

    header_par = section.header.paragraphs[0]
    header_par.text = header_text
    header_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if header_par.runs:
        header_par.runs[0].font.size = Pt(9)
        header_par.runs[0].font.color.rgb = RGBColor(90, 90, 90)

    footer_par = section.footer.paragraphs[0]
    footer_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer_par)


def add_cover_page(doc: Document, title: str, subtitle: str) -> None:
    """Add a clean title page."""
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p1.add_run(title)
    r1.bold = True
    r1.font.size = Pt(24)
    r1.font.color.rgb = RGBColor(26, 58, 99)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(subtitle)
    r2.italic = True
    r2.font.size = Pt(13)
    r2.font.color.rgb = RGBColor(70, 90, 115)

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p3.add_run(f"Generated on {date.today().isoformat()}")
    r3.font.size = Pt(11)
    r3.font.color.rgb = RGBColor(95, 95, 95)

    doc.add_page_break()


def add_code_paragraph(doc: Document, text: str) -> None:
    """Add monospaced code-like paragraph."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(10)


def add_image_from_markdown(doc: Document, notes_dir: Path, alt_text: str, rel_path: str) -> None:
    """Insert image and caption from markdown image syntax."""
    image_path = (notes_dir / rel_path).resolve()
    if not image_path.exists():
        p = doc.add_paragraph(f"[Missing image: {rel_path}]")
        p.runs[0].font.color.rgb = RGBColor(170, 20, 20)
        return

    picture = doc.add_picture(str(image_path), width=Inches(6.0))
    picture_par = doc.paragraphs[-1]
    picture_par.alignment = WD_ALIGN_PARAGRAPH.CENTER

    caption = doc.add_paragraph(f"Figure: {alt_text}")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption.runs:
        caption.runs[0].italic = True
        caption.runs[0].font.size = Pt(9)
        caption.runs[0].font.color.rgb = RGBColor(80, 80, 80)


def add_line(doc: Document, line: str, in_code_block: bool, notes_dir: Path) -> bool:
    """Add a markdown line to DOCX and return updated code-block state."""
    stripped = line.rstrip("\n")

    if stripped.strip() == "```":
        return not in_code_block

    if in_code_block:
        add_code_paragraph(doc, stripped)
        return in_code_block

    match = IMAGE_PATTERN.match(stripped.strip())
    if match:
        alt_text, rel_path = match.group(1), match.group(2)
        add_image_from_markdown(doc, notes_dir, alt_text, rel_path)
        return in_code_block

    if stripped.startswith("# "):
        doc.add_heading(stripped[2:].strip(), level=1)
        return in_code_block
    if stripped.startswith("## "):
        doc.add_heading(stripped[3:].strip(), level=2)
        return in_code_block
    if stripped.startswith("### "):
        doc.add_heading(stripped[4:].strip(), level=3)
        return in_code_block

    if stripped.startswith("- "):
        doc.add_paragraph(stripped[2:].strip(), style="List Bullet")
        return in_code_block

    numbered = stripped.lstrip()
    if len(numbered) > 3 and numbered[0].isdigit() and numbered[1] == "." and numbered[2] == " ":
        doc.add_paragraph(numbered[3:].strip(), style="List Number")
        return in_code_block

    if stripped.startswith("|"):
        add_code_paragraph(doc, stripped)
        return in_code_block

    if stripped.strip() == "":
        doc.add_paragraph("")
        return in_code_block

    doc.add_paragraph(stripped)
    return in_code_block


def append_markdown_to_document(doc: Document, md_path: Path, notes_dir: Path) -> None:
    """Append markdown content to an existing DOCX document."""
    in_code_block = False
    for line in md_path.read_text(encoding="utf-8").splitlines(keepends=True):
        in_code_block = add_line(doc, line, in_code_block, notes_dir)


def extract_title(md_path: Path) -> str:
    """Extract first markdown H1 title."""
    for line in md_path.read_text(encoding="utf-8").splitlines():
        if line.startswith("# "):
            return line[2:].strip()
    return md_path.stem


def convert_markdown_to_docx(md_path: Path, docx_path: Path, notes_dir: Path) -> None:
    """Convert one markdown file to a polished DOCX."""
    doc = Document()
    configure_document(doc)

    note_title = extract_title(md_path)
    add_header_footer(doc, "2D Transformations - Lecture Notes")
    add_cover_page(doc, note_title, "Detailed theory, exact matrices, and Cartesian graph")
    append_markdown_to_document(doc, md_path, notes_dir)

    doc.save(str(docx_path))


def add_simple_toc(doc: Document, notes_dir: Path) -> None:
    """Add a static table of contents section."""
    doc.add_heading("Table of Contents", level=1)
    for index, name in enumerate(SOURCE_FILES, start=1):
        title = extract_title(notes_dir / name)
        doc.add_paragraph(f"{index}. {title}", style="List Number")
    doc.add_page_break()


def main() -> None:
    notes_dir = Path(__file__).resolve().parent

    for name in SOURCE_FILES:
        md_path = notes_dir / name
        docx_path = notes_dir / f"{md_path.stem}.docx"
        convert_markdown_to_docx(md_path, docx_path, notes_dir)

    combined = Document()
    configure_document(combined)
    add_header_footer(combined, "2D Transformations - Combined Note Set")
    add_cover_page(
        combined,
        "2D Transformations - Complete Professional Notes",
        "Translation, Scaling, Rotation, Reflection, Shearing, and Composition",
    )
    add_simple_toc(combined, notes_dir)

    for index, name in enumerate(SOURCE_FILES):
        md_path = notes_dir / name
        append_markdown_to_document(combined, md_path, notes_dir)
        if index != len(SOURCE_FILES) - 1:
            combined.add_page_break()

    combined.save(str(notes_dir / COMBINED_OUTPUT))
    print("DOCX generation complete.")


if __name__ == "__main__":
    main()
